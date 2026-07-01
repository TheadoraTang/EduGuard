"""Prepare anonymized reply comparison for human review.

Joins vanilla and EduGuard replies, randomly assigns A/B labels,
and produces a blinded review file plus an unblinding key.
"""

from __future__ import annotations

import argparse
import json
import random
from pathlib import Path


def read_jsonl(path: Path) -> list[dict]:
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Prepare blinded reply comparison for human review."
    )
    parser.add_argument(
        "--vanilla", type=Path,
        default=Path("data/processed/vanilla_replies.jsonl"),
        help="Output from generate_vanilla_replies.py",
    )
    parser.add_argument(
        "--eduguard", type=Path,
        default=Path("data/processed/eduguard_replies.jsonl"),
        help="Output from generate_eduguard_replies.py",
    )
    parser.add_argument(
        "--original", type=Path,
        default=Path("data/final/studychat_behavior_test.jsonl"),
        help="Original input JSONL for student text and metadata",
    )
    parser.add_argument(
        "--output", type=Path,
        default=Path("data/processed/reply_comparison.jsonl"),
        help="Blinded review output",
    )
    parser.add_argument(
        "--seed", type=int, default=42,
        help="Random seed for reproducible A/B assignment",
    )
    parser.add_argument(
        "--limit", type=int, default=0,
        help="Max comparison pairs to include (0 = all)",
    )
    parser.add_argument(
        "--excel", type=Path, default=None,
        help="Optional Excel output path",
    )
    args = parser.parse_args()

    # Load all sources
    if not args.vanilla.exists():
        raise SystemExit(f"Vanilla replies not found: {args.vanilla}")
    if not args.eduguard.exists():
        raise SystemExit(f"EduGuard replies not found: {args.eduguard}")

    vanilla_rows = read_jsonl(args.vanilla)
    eduguard_rows = read_jsonl(args.eduguard)

    vanilla_map = {r["sample_id"]: r for r in vanilla_rows}
    eduguard_map = {r["sample_id"]: r for r in eduguard_rows}

    # Find common sample_ids
    common = sorted(set(vanilla_map) & set(eduguard_map))
    if not common:
        raise SystemExit(
            f"No common sample_ids found. "
            f"Vanilla: {len(vanilla_map)}, EduGuard: {len(eduguard_map)}"
        )

    if args.limit > 0 and args.limit < len(common):
        common = common[: args.limit]

    # Load original data for student_text and metadata
    original_map: dict[str, dict] = {}
    if args.original.exists():
        for row in read_jsonl(args.original):
            original_map[row["sample_id"]] = row

    # Build comparison rows
    rng = random.Random(args.seed)
    review_rows = []
    unblinding = {
        "seed": args.seed,
        "note": "For each sample_id, a_is_eduguard=true means Reply A is from EduGuard.",
        "assignments": {},
    }

    for idx, sid in enumerate(common, start=1):
        v = vanilla_map[sid]
        e = eduguard_map[sid]
        orig = original_map.get(sid, {})

        swap = rng.random() < 0.5

        reply_a = e["eduguard_reply"] if swap else v["vanilla_reply"]
        reply_b = v["vanilla_reply"] if swap else e["eduguard_reply"]

        review_rows.append({
            "index": idx,
            "sample_id": sid,
            "chat_id_hash": orig.get("chat_id_hash", v.get("chat_id_hash", "")),
            "topic": orig.get("topic", ""),
            "student_text": orig.get("text", ""),
            "reviewed_label": orig.get("reviewed_label", ""),
            "reply_a": reply_a,
            "reply_b": reply_b,
            "preference": "",
            "notes": "",
        })

        unblinding["assignments"][sid] = {
            "a_is_eduguard": swap,
            "eduguard_behavior_type": e.get("behavior_type", ""),
            "eduguard_risk_level": e.get("risk_level", ""),
            "eduguard_strategy_type": e.get("strategy_type", ""),
        }

    # Write blinded review file
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as f:
        for row in review_rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    # Write unblinding key (separate file, NOT for reviewers)
    key_path = args.output.with_suffix(".unblinding.json")
    key_path.write_text(
        json.dumps(unblinding, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    # Summary
    a_eduguard_count = sum(
        1 for v in unblinding["assignments"].values() if v["a_is_eduguard"]
    )
    b_eduguard_count = len(unblinding["assignments"]) - a_eduguard_count

    # Print summary
    summary = {
        "script": "prepare_reply_comparison",
        "total_pairs": len(review_rows),
        "vanilla_count": len(vanilla_map),
        "eduguard_count": len(eduguard_map),
        "common_count": len(common),
        "a_is_eduguard_count": a_eduguard_count,
        "b_is_eduguard_count": b_eduguard_count,
        "blind_ratio": f"{a_eduguard_count}/{len(review_rows)}",
        "jsonl_output": str(args.output),
        "unblinding_key": str(key_path),
    }

    print(json.dumps(summary, ensure_ascii=False, indent=2))

    # Excel generation (if requested and openpyxl available)
    if args.excel:
        try:
            _generate_excel(review_rows, args.excel, summary)
            print(f"Excel workbook written to {args.excel}")
        except ImportError:
            print(
                "openpyxl not installed. Install with: pip install openpyxl\n"
                "Then re-run with: --excel output/spreadsheets/ReplyComparison.xlsx"
            )


def _generate_excel(
    rows: list[dict],
    output_path: Path,
    summary: dict,
) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    # --- Sheet 1: 回复比较 ---
    ws = wb.active
    assert ws is not None
    ws.title = "回复比较"

    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = [
        "序号", "学生提问", "回复A", "回复B",
        "偏好选择\n(A/B/相似/均差)", "备注",
        "人工标签", "主题", "样本ID",
    ]
    col_widths = [6, 40, 50, 50, 12, 25, 18, 8, 30]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    label_fill_map = {
        "answer_seeking": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "concept_understanding": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "debugging_help": PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
        "task_clarification": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "reflection_verification": PatternFill(start_color="E4DFEC", end_color="E4DFEC", fill_type="solid"),
    }

    for row_idx, row in enumerate(rows, start=2):
        values = [
            row["index"],
            row["student_text"],
            row["reply_a"],
            row["reply_b"],
            "",
            "",
            row.get("reviewed_label", ""),
            row.get("topic", ""),
            row["sample_id"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap
            cell.border = thin_border

        # Color label column
        label = row.get("reviewed_label", "")
        if label in label_fill_map:
            ws.cell(row=row_idx, column=7).fill = label_fill_map[label]

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    # --- Sheet 2: 评审指南 ---
    ws2 = wb.create_sheet("评审指南")
    guide = [
        ["EduGuard 回复对比实验 — 人工评审指南", "", ""],
        ["", "", ""],
        ["评价维度", "说明", "评分/判断方式"],
        [
            "教学针对性",
            "回复是否针对学生的具体问题，而非给出通用回答",
            "比较 A 和 B，选择更有针对性的一方",
        ],
        [
            "促进学生思考",
            "回复是引导思考还是直接给答案。好的回复应引导学生自己发现答案",
            "比较 A 和 B，选择更能促进思考的一方",
        ],
        [
            "表达清晰度",
            "回复是否清晰、易懂、结构合理",
            "比较 A 和 B，选择表达更清晰的一方",
        ],
        [
            "是否泄露答案",
            "回复是否直接给出了完整答案或可直接提交的代码",
            "是 / 否 / 部分",
        ],
        [
            "总体偏好",
            "综合以上维度，你更偏好哪个回复",
            "A / B / 相似 / 均差",
        ],
        ["", "", ""],
        ["评审步骤", "", ""],
        ["1. 阅读学生提问，理解学生遇到的问题", "", ""],
        ["2. 分别阅读回复A和回复B", "", ""],
        ["3. 比较两个回复在各维度上的表现", "", ""],
        ["4. 在Excel的 回复比较 工作表中填写偏好选择和备注", "", ""],
        ["", "", ""],
        ["注意", "", ""],
        ["- A和B的顺序是随机的，不要根据位置判断优劣", "", ""],
        ["- 如果两个回复质量相当，可以选择 相似", "", ""],
        ["- 如果两个回复质量都差，可以选择 均差 并说明原因", "", ""],
        ["- 备注中可以记录你的观察、判断理由或特殊注意事项", "", ""],
    ]
    for row_idx, row_data in enumerate(guide, start=1):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            if row_idx in (1, 4, 15):
                cell.font = Font(bold=True)
            if row_idx == 1:
                cell.font = Font(bold=True, size=14)

    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 55
    ws2.column_dimensions["C"].width = 30

    # --- Sheet 3: 统计摘要 ---
    ws3 = wb.create_sheet("统计摘要")
    stats = [
        ["EduGuard 回复对比实验 — 统计摘要", "", ""],
        ["", "", ""],
        ["指标", "值", ""],
        ["比较对总数", summary["total_pairs"], ""],
        ["EduGuard 为回复A的数量", summary["a_is_eduguard_count"], ""],
        ["EduGuard 为回复B的数量", summary["b_is_eduguard_count"], ""],
        [
            "盲化比例",
            f"{summary['a_is_eduguard_count']}/{summary['total_pairs']}",
            "",
        ],
    ]
    for row_idx, row_data in enumerate(stats, start=1):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=value)
            if row_idx == 1:
                cell.font = Font(bold=True, size=14)

    # Label distribution
    label_counts: dict[str, int] = {}
    for row in rows:
        label = row.get("reviewed_label", "unknown")
        label_counts[label] = label_counts.get(label, 0) + 1

    start_row = len(stats) + 2
    ws3.cell(row=start_row, column=1, value="标签分布").font = Font(bold=True)
    for i, (label, count) in enumerate(sorted(label_counts.items()), start=1):
        ws3.cell(row=start_row + i, column=1, value=label)
        ws3.cell(row=start_row + i, column=2, value=count)

    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 20

    wb.save(output_path)


if __name__ == "__main__":
    main()
