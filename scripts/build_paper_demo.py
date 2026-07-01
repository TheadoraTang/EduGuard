from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "template" / "论文模板.docx"
OUT = ROOT / "docs" / "EduGuard论文初稿.docx"


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def set_page(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)


def set_columns(section, count: int = 2) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    cols_el = cols[0] if cols else OxmlElement("w:cols")
    cols_el.set(qn("w:num"), str(count))
    cols_el.set(qn("w:space"), "425")
    if not cols:
        sect_pr.append(cols_el)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in kwargs.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def set_cell_text(cell, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(8)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_three_line_table(doc: Document, title_cn: str, title_en: str, headers, rows) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title_cn + "\n")
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(9)
    r2 = p.add_run(title_en)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    remove_table_borders(table)

    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 and len(str(value)) > 16 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[i], str(value), align=align)

    top = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    mid = {"val": "single", "sz": "6", "space": "0", "color": "000000"}
    bottom = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    for cell in table.rows[0].cells:
        set_cell_border(cell, top=top, bottom=mid)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom=bottom)

    doc.add_paragraph()


def set_run_font(run, size=10, bold=False, italic=False, name="宋体") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def para(doc: Document, text: str = "", align=None, size=10, bold=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.first_line_indent = Pt(20) if text else None
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


def add_citation(paragraph, num: int) -> None:
    run = paragraph.add_run(f"[{num}]")
    set_run_font(run, size=8, name="Times New Roman")
    run.font.superscript = True


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=11.5 if level <= 2 else 10.5, bold=True, name="黑体")


def add_ref(doc: Document, n: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    r = p.add_run(f"[{n}] {text}")
    set_run_font(r, size=8, name="宋体")


def build() -> None:
    shutil.copyfile(TEMPLATE, OUT)
    doc = Document(str(OUT))
    clear_body(doc)
    for section in doc.sections:
        set_page(section)
        set_columns(section, 1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("EduGuard：面向学生-AI交互的行为感知支架式教学系统")
    set_run_font(r, size=18, bold=True, name="黑体")

    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = authors.add_run("作者名")
    set_run_font(r, size=10.5, name="楷体")

    aff = doc.add_paragraph()
    aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = aff.add_run("单位全名 学院(系)全名, 省 市 邮政编码")
    set_run_font(r, size=9, name="楷体")

    p = para(doc, "摘  要：学生使用生成式人工智能辅助学习时，其交互行为具有明显差异。部分学生希望理解概念，部分学生带着错误信息寻求调试帮助，也有学生直接索要可提交答案。通用大语言模型通常对这些请求采用相似的直接回答方式，难以区分健康学习行为与浅层学习风险。针对这一问题，设计了 EduGuard 行为感知支架式教学系统。系统首先对学生消息进行行为分类，然后基于证据规则评估浅层学习风险，再选择支架策略，并通过自然教学 persona 生成回复。实验结果表明，TF-IDF 与逻辑回归分类器在 StudyChat 五分类任务上达到 70.3% 的准确率，优于规则方法和 LLM prompt 基线；在回复质量对比实验中，EduGuard v3 在教学引导性、答案保护性和启发性上分别取得 74.0%、78.7% 和 66.0% 的相对胜率。结果说明，行为感知管线能够在一定程度上降低直接给答案的风险，并把通用 LLM 转化为更具有教学支架意识的学习助手。", size=10)
    p.paragraph_format.first_line_indent = None
    p = para(doc, "关键词：生成式人工智能；智能教学系统；学习支架；行为分类；学习分析", size=10)
    p.paragraph_format.first_line_indent = None

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EduGuard: A behavior-aware scaffolded tutoring system for student-AI interaction")
    set_run_font(r, size=14, bold=True, name="Times New Roman")
    p = para(doc, "Abstract: When students use generative AI for learning, their interaction patterns vary substantially. Some students seek conceptual understanding, some ask for debugging help with concrete errors, while others directly request complete answers. General-purpose large language models often respond to these requests in a similar way and therefore fail to distinguish productive help seeking from shallow learning risk. This paper presents EduGuard, a behavior-aware scaffolded tutoring system. The system classifies student behavior, estimates shallow learning risk through evidence-based rules, selects a scaffolding strategy, and generates a response through a natural pedagogical persona. Experiments show that a TF-IDF plus logistic regression classifier achieves 70.3% accuracy on the StudyChat five-class behavior classification task. In response quality evaluation, EduGuard v3 outperforms a vanilla LLM on pedagogical guidance, answer protection, and inspiration. These results suggest that behavior-aware scaffolding can help transform LLM responses from direct answer delivery to guided learning support.", size=10, )
    p.paragraph_format.first_line_indent = None
    p = para(doc, "Key words: generative AI; intelligent tutoring system; scaffolding; behavior classification; learning analytics", size=10)
    p.paragraph_format.first_line_indent = None

    # Main text starts in two columns.
    doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(doc.sections[-1], 2)

    p = para(doc, "生成式人工智能正在快速进入学习场景。学生可以通过自然语言获得解释、代码、调试建议和写作辅助，这为个性化学习提供了新的可能。然而，通用大语言模型并不天然具备教学判断能力：当学生直接索要答案时，模型可能给出完整解法；当学生需要调试训练时，模型也可能直接替代其完成排查过程。这种便利性可能削弱学生的主动思考和问题分解能力。学习支架理论强调，教学支持应当根据学习者当前能力与任务难度进行适配，并逐步引导学习者完成其独立状态下尚不能完成的任务")
    add_citation(p, 1)
    p.add_run("。智能教学系统研究也表明，系统若能建模学生状态并提供个性化反馈，能够接近甚至部分替代一对一辅导的效果")
    add_citation(p, 6)
    p.add_run("。")

    para(doc, "本文提出 EduGuard，一个面向学生-AI交互的行为感知支架式教学系统。与直接把 LLM 包装成聊天助手不同，EduGuard 在回复生成前先分析学生如何使用 AI，判断是否存在浅层学习风险，然后选择相应的支架策略。系统的目标不是禁止学生使用 AI，而是把学生与 AI 的交互从“得到答案”转化为“得到引导”。")

    heading(doc, "1  相关工作", 1)
    para(doc, "学习支架最早强调教师在问题解决过程中对学习者提供临时性支持，帮助其完成更高水平的认知活动。后续研究进一步指出，技术环境中的支架需要持续诊断学习者状态，并根据任务进展调整支持方式")
    add_citation(doc.paragraphs[-1], 3)
    doc.paragraphs[-1].add_run("。在人工智能教育系统中，Cognitive Tutor、AutoTutor 等系统通过学生模型、反馈模型和自然语言对话提供个性化支持")
    add_citation(doc.paragraphs[-1], 7)
    add_citation(doc.paragraphs[-1], 15)
    doc.paragraphs[-1].add_run("。")
    para(doc, "大语言模型为自然语言教学系统提供了更强的生成能力，但也带来新的教育风险。已有研究总结了 LLM 在教育中的机会与挑战，包括反馈生成、内容生成、隐私、透明性和过度依赖等问题")
    add_citation(doc.paragraphs[-1], 13)
    add_citation(doc.paragraphs[-1], 14)
    doc.paragraphs[-1].add_run("。因此，EduGuard 的研究重点不是单纯提升回答能力，而是让 LLM 回复受到学习行为分析和支架策略的约束。")

    heading(doc, "2  EduGuard 管线设计", 1)
    para(doc, "EduGuard 的核心管线包括四个阶段：行为分类、风险评估、支架策略选择和教学回复生成。第一阶段识别学生消息属于概念理解、答案索取、调试帮助、任务澄清或反思验证中的哪一类；第二阶段基于行为类型、是否提供个人尝试、是否反复索要答案等证据计算风险分数；第三阶段根据行为类型和风险等级选择支架策略；第四阶段将策略转化为自然教学 persona，并由 LLM 生成最终回复。")

    add_three_line_table(
        doc,
        "表1  学生行为标签定义",
        "Table 1  Definitions of student behavior labels",
        ["标签", "含义", "典型意图"],
        [
            ["answer_seeking", "索要完整答案、代码或可提交成品", "直接完成任务"],
            ["concept_understanding", "请求概念解释、原理或思路", "理解知识"],
            ["debugging_help", "提供代码或错误并请求定位修复", "排查问题"],
            ["task_clarification", "澄清题目要求、输入输出或约束", "明确任务"],
            ["reflection_verification", "检查、反思或质疑已有结果", "验证理解"],
        ],
    )

    para(doc, "风险评估采用加权证据规则。答案索取、重复索要和要求可提交代码会提高风险分数；提供个人尝试、错误信息或反思性表述会降低风险分数。这种设计使系统能够区分“需要帮助的学习者”和“希望跳过学习过程的学习者”。在高风险场景下，系统优先采用 guided_questions 策略；在调试场景下，系统采用 debugging_scaffold 策略；在反思验证场景下，系统采用 verification_prompt 策略。")

    heading(doc, "3  数据集与标注", 1)
    para(doc, "实验使用 StudyChat 和 MathDial 两类数据。StudyChat 用于学生行为分类任务，最终保留 14105 条人工复核样本，并按完整对话划分训练集、开发集和测试集，避免同一对话跨集合泄漏。MathDial 用于支架策略与教师反馈分析，共包含 1977 个教师回复轮次，人工评分显示其教师回复整体具有较高的针对性、促进思考和清晰度。")
    add_three_line_table(
        doc,
        "表2  StudyChat 数据集分布",
        "Table 2  Distribution of the StudyChat dataset",
        ["划分", "样本数"],
        [["train", "10026"], ["dev", "2141"], ["test", "1938"], ["total", "14105"]],
    )

    heading(doc, "4  实验一：行为分类", 1)
    para(doc, "为了验证管线第一阶段的可靠性，比较了规则分类器、LLM zero-shot、LLM few-shot 和 TF-IDF+LR 四种方法。TF-IDF+LR 使用 word 1-2 gram 与 char 3-5 gram 特征，并设置 class_weight=balanced，以缓解类别不均衡问题。")
    add_three_line_table(
        doc,
        "表3  行为分类实验结果",
        "Table 3  Results of behavior classification",
        ["方法", "Accuracy", "Macro-F1", "Weighted-F1"],
        [
            ["规则分类器", "50.6%", "0.342", "0.459"],
            ["LLM Zero-shot", "27.6%", "0.307", "0.288"],
            ["LLM Few-shot", "56.6%", "0.493", "0.605"],
            ["TF-IDF+LR", "70.3%", "0.622", "0.707"],
        ],
    )
    para(doc, "结果表明，轻量机器学习分类器在该任务上明显优于直接提示 LLM。LLM zero-shot 容易偏向少数策略标签，导致策略错配；few-shot 能明显改善，但仍低于 TF-IDF+LR。该结果说明，在教学管线中，稳定、可复现的行为分类器比单纯依赖大模型即时判断更适合作为前置模块。")

    heading(doc, "5  实验二：回复质量对比", 1)
    para(doc, "回复质量实验比较普通 LLM 与 EduGuard 管线。早期版本直接把风险分数和策略标签注入 system prompt，导致回复过于机械；替换分类器后胜率提升；最终 v3 将策略转化为自然教学 persona，使模型以更自然的苏格拉底式导师风格进行反馈。")
    add_three_line_table(
        doc,
        "表4  EduGuard v3 与通用 LLM 的多维度比较",
        "Table 4  Multi-dimensional comparison between EduGuard v3 and vanilla LLM",
        ["维度", "EduGuard", "Vanilla", "相似", "Edu 胜率"],
        [
            ["教学引导性", "37", "13", "0", "74.0%"],
            ["答案保护性", "37", "10", "3", "78.7%"],
            ["启发性", "33", "17", "0", "66.0%"],
            ["针对性", "15", "34", "1", "30.6%"],
            ["总体偏好", "23", "25", "2", "47.9%"],
        ],
    )
    para(doc, "EduGuard 在教学引导性、答案保护性和启发性上明显优于通用 LLM，说明行为感知支架能够减少直接泄露答案，并鼓励学生继续思考。但针对性结果较低，说明当前 persona 在保护答案和精准回应之间仍存在取舍。后续需要在低风险场景适度增加实质性解释，并在调试场景中更直接地分析学生已经提供的代码和错误。")

    heading(doc, "6  系统实现", 1)
    para(doc, "EduGuard 被实现为一个双角色 Web 教学平台。学生端提供对话界面、任务选择、作业提交和学习反思；教师端提供风险分布、行为分布、对话诊断和人工审核。后端使用 FastAPI 与 SQLAlchemy，前端使用 React、TypeScript、Ant Design 和 ECharts。系统在学生发送消息后执行完整管线，并将行为分类、风险证据、支架策略和最终回复持久化，便于教师进行 Human-in-the-Loop 审核。")

    heading(doc, "7  结束语", 1)
    para(doc, "本文提出并实现了 EduGuard 行为感知支架式教学系统。系统通过行为分类、风险评估、支架策略选择和教学回复生成，把通用 LLM 转化为更具教学支架意识的学习助手。实验结果说明，稳定的行为分类器是管线有效性的基础，而自然教学 persona 是改善回复质量的关键。当前系统仍存在分类边界模糊、回复针对性不足和评审样本规模有限等问题。未来工作将进一步引入对话历史特征、尝试更强的分类模型，并开展多标注者一致性检验。")

    heading(doc, "参考文献：", 1)
    refs = [
        "WOOD D, BRUNER J S, ROSS G. The role of tutoring in problem solving[J]. Journal of Child Psychology and Psychiatry, 1976, 17(2): 89-100.",
        "VYGOTSKY L S. Mind in society: The development of higher psychological processes[M]. Cambridge: Harvard University Press, 1978.",
        "PUNTAMBEKAR S, HUBSCHER R. Tools for scaffolding students in a complex learning environment: What have we gained and what have we missed?[J]. Educational Psychologist, 2005, 40(1): 1-12.",
        "CHI M T H, SILER S A, JEONG H, et al. Learning from human tutoring[J]. Cognitive Science, 2001, 25(4): 471-533.",
        "CHI M T H, WYLIE R. The ICAP framework: Linking cognitive engagement to active learning outcomes[J]. Educational Psychologist, 2014, 49(4): 219-243.",
        "VANLEHN K. The relative effectiveness of human tutoring, intelligent tutoring systems, and other tutoring systems[J]. Educational Psychologist, 2011, 46(4): 197-221.",
        "CORBETT A T, KOEDINGER K R, ANDERSON J R. Intelligent tutoring systems[M]//HELANDER M G, LANDAUER T K, PRABHU P V. Handbook of human-computer interaction. Amsterdam: Elsevier, 1997: 849-874.",
        "KOEDINGER K R, CORBETT A T. Cognitive tutors: Technology bringing learning sciences to the classroom[M]//SAWYER R K. The Cambridge handbook of the learning sciences. Cambridge: Cambridge University Press, 2006.",
        "ALEVEN V, KOEDINGER K R. An effective metacognitive strategy: Learning by doing and explaining with a computer-based Cognitive Tutor[J]. Cognitive Science, 2002, 26(2): 147-179.",
        "ROLL I, ALEVEN V, MCLAREN B M, et al. Improving students' help-seeking skills using metacognitive feedback in an intelligent tutoring system[J]. Learning and Instruction, 2011, 21(2): 267-280.",
        "SHUTE V J. Focus on formative feedback[J]. Review of Educational Research, 2008, 78(1): 153-189.",
        "ZAWACKI-RICHTER O, MARIN V I, BOND M, et al. Systematic review of research on artificial intelligence applications in higher education[J]. International Journal of Educational Technology in Higher Education, 2019, 16: 39.",
        "KASNECI E, SESSLER K, KUCHEMANN S, et al. ChatGPT for good? On opportunities and challenges of large language models for education[J]. Learning and Individual Differences, 2023, 103: 102274.",
        "YAN L, SHA L, ZHAO L, et al. Practical and ethical challenges of large language models in education: A systematic scoping review[EB/OL]. arXiv, 2023[2026-06-25]. https://arxiv.org/abs/2303.13379.",
        "NYE B D, GRAESSER A C, HU X. AutoTutor and family: A review of 17 years of natural language tutoring[J]. International Journal of Artificial Intelligence in Education, 2014, 24: 427-469.",
    ]
    for i, ref in enumerate(refs, start=1):
        add_ref(doc, i, ref)

    doc.save(str(OUT))


if __name__ == "__main__":
    build()
