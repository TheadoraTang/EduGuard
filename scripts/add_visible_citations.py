from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "docs" / "EduGuard论文初稿.docx"
BACKUP = ROOT / "docs" / "EduGuard论文初稿.before-citations.docx"


def set_citation_font(run) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(8)
    run.font.superscript = True


def paragraph_has_all(text: str, citations: list[str]) -> bool:
    return all(c in text for c in citations)


def append_citations(paragraph, citations: list[str]) -> None:
    text = paragraph.text
    if paragraph_has_all(text, citations):
        return

    # Put citations before the final sentence punctuation when possible.
    final_punctuation = ""
    if paragraph.runs and paragraph.runs[-1].text.endswith(("。", ".")):
        final_punctuation = paragraph.runs[-1].text[-1]
        paragraph.runs[-1].text = paragraph.runs[-1].text[:-1]

    for citation in citations:
        if citation not in paragraph.text:
            run = paragraph.add_run(citation)
            set_citation_font(run)

    if final_punctuation:
        paragraph.add_run(final_punctuation)


def main() -> None:
    if not TARGET.exists():
        raise FileNotFoundError(TARGET)

    shutil.copyfile(TARGET, BACKUP)
    doc = Document(str(TARGET))

    citation_plan = [
        (
            "EduGuard 在回复生成前先分析学生如何使用 AI",
            ["[10]", "[11]", "[13]"],
        ),
        (
            "EduGuard 的核心管线包括四个阶段",
            ["[6]", "[7]", "[15]"],
        ),
        (
            "风险评估采用加权证据规则",
            ["[1]", "[3]", "[10]", "[11]"],
        ),
        (
            "实验使用 StudyChat 和 MathDial 两类数据",
            ["[12]", "[13]"],
        ),
        (
            "为了验证管线第一阶段的可靠性",
            ["[12]", "[14]"],
        ),
        (
            "稳定、可复现的行为分类器比单纯依赖大模型即时判断更适合作为前置模块",
            ["[13]", "[14]"],
        ),
        (
            "最终 v3 将策略转化为自然教学 persona",
            ["[4]", "[5]", "[15]"],
        ),
        (
            "说明行为感知支架能够减少直接泄露答案",
            ["[5]", "[10]", "[11]"],
        ),
        (
            "便于教师进行 Human-in-the-Loop 审核",
            ["[6]", "[7]", "[12]"],
        ),
        (
            "本文提出并实现了 EduGuard 行为感知支架式教学系统",
            ["[1]", "[6]", "[13]"],
        ),
    ]

    for paragraph in doc.paragraphs:
        for marker, citations in citation_plan:
            if marker in paragraph.text:
                append_citations(paragraph, citations)

    doc.save(str(TARGET))


if __name__ == "__main__":
    main()
